"""Módulo para geração de relatórios e documentos"""

from pathlib import Path
from datetime import datetime
from typing import Dict, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors

from src.database.storage import DataStorage


class ReportGenerator:
    """Gera relatórios e documentos de risco"""

    def __init__(self, storage: DataStorage):
        self.storage = storage

    def generate_pgr(self, file_path: str, company_info: Dict, scope: str):
        """Gera Programa de Gestão de Riscos em DOCX"""
        doc = Document()

        # Title
        title = doc.add_heading('Programa de Gestão de Riscos (PGR)', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Company info
        doc.add_heading('1. Informações da Empresa', 1)
        doc.add_paragraph(f"Razão Social: {company_info.get('name', 'N/A')}")
        doc.add_paragraph(f"CNPJ: {company_info.get('cnpj', 'N/A')}")
        doc.add_paragraph(f"Endereço: {company_info.get('address', 'N/A')}")
        doc.add_paragraph(f"Responsável: {company_info.get('responsible', 'N/A')}")
        doc.add_paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y')}")

        # Scope
        doc.add_heading('2. Escopo da Avaliação', 1)
        doc.add_paragraph(scope)

        # Methodology
        doc.add_heading('3. Metodologia', 1)
        doc.add_paragraph(
            "A avaliação de riscos foi realizada conforme a norma ISO 31010 - "
            "Gestão de Riscos - Técnicas de Avaliação de Riscos, em conformidade com "
            "as Normas Regulamentadoras (NRs) brasileiras."
        )

        # Environments
        doc.add_heading('4. Ambientes Avaliados', 1)
        environments = self.storage.get_all_environments()

        if environments:
            for env in environments:
                doc.add_heading(f"4.1 {env.name}", 2)
                doc.add_paragraph(f"Tipo: {env.type}")
                doc.add_paragraph(f"Localização: {env.location}")
                doc.add_paragraph(f"Número de Trabalhadores: {env.workers_count}")
                doc.add_paragraph(f"Descrição: {env.description}")
        else:
            doc.add_paragraph("Nenhum ambiente registrado.")

        # Hazards
        doc.add_heading('5. Perigos Identificados', 1)
        hazards = self.storage.get_all_hazards()

        if hazards:
            for hazard in hazards:
                doc.add_heading(f"5.1 {hazard.category.value}", 2)
                doc.add_paragraph(f"Descrição: {hazard.description}")
                doc.add_paragraph(f"Localização: {hazard.location}")
                doc.add_paragraph(f"Pessoas Expostas: {hazard.exposed_people}")
        else:
            doc.add_paragraph("Nenhum perigo identificado.")

        # Risk Assessments
        doc.add_heading('6. Avaliação de Riscos', 1)
        assessments = self.storage.get_all_assessments()

        if assessments:
            # Summary table
            summary_data = [['Nível de Risco', 'Quantidade']]
            risk_counts = {}

            for assess in assessments:
                level = assess.get_risk_level()
                risk_counts[level] = risk_counts.get(level, 0) + 1

            for level, count in sorted(risk_counts.items()):
                summary_data.append([level, str(count)])

            table = doc.add_table(rows=len(summary_data), cols=2)
            table.style = 'Light Grid Accent 1'

            for i, row_data in enumerate(summary_data):
                row = table.rows[i]
                for j, cell_data in enumerate(row_data):
                    row.cells[j].text = cell_data

            doc.add_paragraph()

            # Detailed assessments
            for assess in assessments:
                hazard = self.storage.get_hazard(assess.hazard_id)
                hazard_desc = hazard.description if hazard else "N/A"

                doc.add_heading(f"6.1 {hazard_desc}", 2)
                doc.add_paragraph(f"Nível de Risco: {assess.get_risk_level()}")
                doc.add_paragraph(f"Severidade: {assess.severity.value}")
                doc.add_paragraph(f"Probabilidade: {assess.probability.value}")
                doc.add_paragraph(f"Medidas de Controle: {assess.control_measures}")

                if assess.residual_severity:
                    doc.add_paragraph(f"Severidade Residual: {assess.residual_severity.value}")
                if assess.residual_probability:
                    doc.add_paragraph(f"Probabilidade Residual: {assess.residual_probability.value}")

                if assess.responsible:
                    doc.add_paragraph(f"Responsável: {assess.responsible}")
                if assess.deadline:
                    doc.add_paragraph(f"Prazo: {assess.deadline}")

                doc.add_paragraph()
        else:
            doc.add_paragraph("Nenhuma avaliação de risco realizada.")

        # Conclusions
        doc.add_heading('7. Conclusões e Recomendações', 1)
        doc.add_paragraph(
            "O Programa de Gestão de Riscos deve ser continuamente atualizado e "
            "revisado com base em mudanças no ambiente de trabalho, novas atividades "
            "ou novos riscos identificados."
        )

        doc.save(file_path)

    def generate_pdf_report(self, file_path: str, company_info: Dict, scope: str):
        """Gera Relatório de Análise de Riscos em PDF"""
        doc = SimpleDocTemplate(file_path, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1a1a1a'),
            spaceAfter=30,
            alignment=1  # Center
        )
        story.append(Paragraph('Relatório de Análise de Riscos', title_style))
        story.append(Paragraph(f'ISO 31010 / NRs Brasileiras', styles['Normal']))
        story.append(Spacer(1, 0.3 * inch))

        # Company Info
        story.append(Paragraph('Informações da Empresa', styles['Heading2']))
        company_data = [
            ['Razão Social:', company_info.get('name', 'N/A')],
            ['CNPJ:', company_info.get('cnpj', 'N/A')],
            ['Endereço:', company_info.get('address', 'N/A')],
            ['Responsável:', company_info.get('responsible', 'N/A')],
            ['Data:', datetime.now().strftime('%d/%m/%Y')],
        ]

        company_table = Table(company_data, colWidths=[2 * inch, 4 * inch])
        company_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))

        story.append(company_table)
        story.append(Spacer(1, 0.3 * inch))

        # Scope
        story.append(Paragraph('Escopo da Avaliação', styles['Heading2']))
        story.append(Paragraph(scope, styles['Normal']))
        story.append(Spacer(1, 0.2 * inch))

        # Statistics
        stats = self.storage.get_statistics()
        story.append(Paragraph('Resumo Executivo', styles['Heading2']))

        summary_data = [
            ['Métrica', 'Valor'],
            ['Ambientes Avaliados', str(stats['total_environments'])],
            ['Perigos Identificados', str(stats['total_hazards'])],
            ['Avaliações de Risco', str(stats['total_assessments'])],
        ]

        summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))

        story.append(summary_table)
        story.append(Spacer(1, 0.2 * inch))

        # Risk Distribution
        if stats['risks_by_level']:
            story.append(Paragraph('Distribuição de Riscos', styles['Heading2']))

            risk_data = [['Nível de Risco', 'Quantidade']]
            for level, count in sorted(stats['risks_by_level'].items()):
                risk_data.append([level, str(count)])

            risk_table = Table(risk_data, colWidths=[3 * inch, 2 * inch])
            risk_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))

            story.append(risk_table)
            story.append(Spacer(1, 0.2 * inch))

        story.append(PageBreak())

        # Build PDF
        doc.build(story)
